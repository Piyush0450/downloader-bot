import asyncio
import os
import shutil
import subprocess
import zipfile
import csv
import json
from pathlib import Path
from PIL import Image
from pypdf import PdfReader
from pdf2docx import Converter as PDF2DOCXConverter
import docx2txt
import mammoth
import openpyxl

from config.settings import CONVERTED_DIR
from utils.helpers import generate_temp_path
from utils.logger import logger

def get_cmd_path(name: str) -> str:
    path = shutil.which(name)
    if path:
        return path
    if os.name == "nt":
        if name in ("7z", "7za"):
            for p in (Path("C:/Program Files/7-Zip/7z.exe"), Path("C:/Program Files (x86)/7-Zip/7z.exe")):
                if p.exists():
                    return str(p)
        elif name in ("libreoffice", "soffice"):
            for p in (Path("C:/Program Files/LibreOffice/program/soffice.exe"), Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe")):
                if p.exists():
                    return str(p)
    return None

# --- SPREADSHEETS PURE PYTHON CONVERTERS ---

def convert_xlsx_to_json_sync(in_path: Path, out_path: Path):
    wb = openpyxl.load_workbook(in_path, read_only=True, data_only=True)
    result = {}
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sheet_data = []
        for row in sheet.iter_rows(values_only=True):
            if any(row is not None for row in row):
                sheet_data.append(list(row))
        result[sheet_name] = sheet_data
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=4, ensure_ascii=False)

def convert_csv_to_json_sync(in_path: Path, out_path: Path):
    result = []
    with open(in_path, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            result.append(row)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=4, ensure_ascii=False)

def convert_csv_to_xlsx_sync(in_path: Path, out_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    with open(in_path, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            ws.append(row)
    wb.save(out_path)

def convert_xlsx_to_csv_sync(in_path: Path, out_path: Path):
    wb = openpyxl.load_workbook(in_path, read_only=True, data_only=True)
    sheet = wb.active
    with open(out_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(list(row))

# --- ARCHIVES EXTRACTOR ---

def extract_archive_sync(in_path: Path, out_zip_path: Path) -> bool:
    """Extract zip/rar/7z archives using 7z command or python zipfile."""
    temp_extract_dir = in_path.parent / f"extracted_{in_path.stem}"
    temp_extract_dir.mkdir(exist_ok=True)
    
    seven_z = get_cmd_path("7z")
    success = False
    
    try:
        if seven_z:
            # 7z x input -oextracted_dir -y
            cmd = [seven_z, "x", str(in_path), f"-o{temp_extract_dir}", "-y"]
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            success = True
        elif in_path.suffix.lower() == ".zip":
            # Fallback for zip files
            with zipfile.ZipFile(in_path, 'r') as zip_ref:
                zip_ref.extractall(temp_extract_dir)
            success = True
        else:
            raise FileNotFoundError("System requires '7z' CLI utility to extract RAR/7Z files.")
            
        # Re-archive the folder contents into a single safe zipfile
        files = list(temp_extract_dir.rglob("*"))
        if files:
            with zipfile.ZipFile(out_zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for file in files:
                    if file.is_file():
                        # Calculate relative path
                        rel_path = file.relative_to(temp_extract_dir)
                        zipf.write(file, rel_path)
            success = True
            
    except Exception as e:
        logger.error(f"Archive extraction failed: {e}")
        success = False
    finally:
        shutil.rmtree(temp_extract_dir, ignore_errors=True)
        
    return success

# --- SYSTEM CONVERTERS ---

def convert_image_sync(in_path: Path, out_path: Path, target_ext: str):
    with Image.open(in_path) as img:
        target_ext = target_ext.lower().replace(".", "")
        if target_ext in ("jpg", "jpeg"):
            if img.mode in ("RGBA", "LA", "P"):
                background = Image.new("RGB", img.size, (255, 255, 255))
                background.paste(img, mask=img.split()[3] if img.mode == "RGBA" else None)
                img = background
            else:
                img = img.convert("RGB")
            img.save(out_path, "JPEG", quality=90)
        elif target_ext == "pdf":
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
            img.save(out_path, "PDF", resolution=100.0)
        elif target_ext == "ico":
            img.save(out_path, "ICO", sizes=[(32, 32), (64, 64)])
        else:
            img.save(out_path, target_ext.upper())

def convert_document_libreoffice_sync(in_path: Path, target_format: str, out_dir: Path) -> Path:
    libreoffice = get_cmd_path("libreoffice") or get_cmd_path("soffice")
    if not libreoffice:
        raise FileNotFoundError("LibreOffice CLI utility not found.")
        
    cmd = [
        libreoffice,
        "--headless",
        "--convert-to",
        target_format,
        "--outdir",
        str(out_dir),
        str(in_path)
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    
    # Locate output file
    expected_filename = f"{in_path.stem}.{target_format}"
    out_path = out_dir / expected_filename
    if out_path.exists():
        return out_path
        
    for file in out_dir.iterdir():
        if file.stem == in_path.stem and file.suffix.lstrip(".").lower() == target_format.lower():
            return file
            
    raise FileNotFoundError(f"LibreOffice conversion failed to write: {expected_filename}")

def convert_media_ffmpeg_sync(in_path: Path, out_path: Path, target_ext: str):
    ffmpeg = get_cmd_path("ffmpeg")
    if not ffmpeg:
        raise FileNotFoundError("FFmpeg CLI utility not found.")
        
    cmd = [ffmpeg, "-y", "-i", str(in_path)]
    target_ext = target_ext.lower().replace(".", "")
    
    if target_ext in ("mp3", "wav", "ogg", "aac", "flac"):
        # Audio output specs
        cmd.append("-vn")
        if target_ext == "mp3":
            cmd.extend(["-q:a", "2"])
        elif target_ext == "flac":
            cmd.extend(["-compression_level", "5"])
        cmd.append(str(out_path))
    elif target_ext == "gif":
        cmd.extend([
            "-vf", "fps=10,scale=480:-1:flags=lanczos,split[s0][s1];[s0]palettegen[p];[s1][p]paletteuse",
            "-loop", "0",
            str(out_path)
        ])
    else:
        # Standard video transcoding
        cmd.extend([
            "-c:v", "libx264",
            "-c:a", "aac",
            "-strict", "experimental",
            str(out_path)
        ])
        
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

# --- PDF PAGE RENDERING ---
def convert_pdf_to_images_sync(in_path: Path, out_zip_path: Path) -> bool:
    pdftoppm = get_cmd_path("pdftoppm")
    temp_pages_dir = in_path.parent / f"pdf_pages_{in_path.stem}"
    temp_pages_dir.mkdir(exist_ok=True)
    success = False
    try:
        if pdftoppm:
            cmd = [pdftoppm, "-png", "-r", "150", str(in_path), str(temp_pages_dir / "page")]
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            success = True
        else:
            reader = PdfReader(in_path)
            count = 0
            for page in reader.pages:
                for img_obj in page.images:
                    count += 1
                    with open(temp_pages_dir / f"page_{count}.png", "wb") as f:
                        f.write(img_obj.data)
            if count > 0:
                success = True
                
        files = list(temp_pages_dir.glob("*"))
        if files:
            with zipfile.ZipFile(out_zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for file in files:
                    zipf.write(file, file.name)
            success = True
    except Exception as e:
        logger.error(f"PDF to images conversion failed: {e}")
    finally:
        shutil.rmtree(temp_pages_dir, ignore_errors=True)
    return success

# --- MAIN ASYNC EXECUTOR ROUTER ---

async def execute_format_conversion(in_path: Path, target_ext: str) -> Path:
    """Orchestrates conversion pipeline based on formats and system libraries."""
    source_ext = in_path.suffix.lstrip(".").lower()
    target_ext = target_ext.lower().replace(".", "")
    out_path = CONVERTED_DIR / f"conv_{generate_temp_path(target_ext).name}"
    
    # 1. Archive extraction
    if source_ext in ("zip", "rar", "7z") and target_ext == "extract":
        zip_out = CONVERTED_DIR / f"extracted_{generate_temp_path('zip').name}"
        success = await asyncio.to_thread(extract_archive_sync, in_path, zip_out)
        if success and zip_out.exists():
            return zip_out
        raise RuntimeError("Archive extraction failed.")
        
    # 2. Image conversions
    if source_ext in ("png", "jpg", "jpeg", "webp", "bmp", "tiff"):
        await asyncio.to_thread(convert_image_sync, in_path, out_path, target_ext)
        return out_path
        
    # 3. Spreadsheets (Pure python xlsx/csv/json)
    if source_ext == "xlsx" and target_ext == "json":
        await asyncio.to_thread(convert_xlsx_to_json_sync, in_path, out_path)
        return out_path
    if source_ext == "csv" and target_ext == "json":
        await asyncio.to_thread(convert_csv_to_json_sync, in_path, out_path)
        return out_path
    if source_ext == "csv" and target_ext == "xlsx":
        await asyncio.to_thread(convert_csv_to_xlsx_sync, in_path, out_path)
        return out_path
    if source_ext == "xlsx" and target_ext == "csv":
        await asyncio.to_thread(convert_xlsx_to_csv_sync, in_path, out_path)
        return out_path
        
    # 4. PDF extractions
    if source_ext == "pdf":
        if target_ext == "docx":
            # PDF2DOCX
            await asyncio.to_thread(lambda: PDF2DOCXConverter(str(in_path)).convert(str(out_path)) or None)
            return out_path
        elif target_ext == "txt":
            # Extract text
            def extract_pdf_txt():
                reader = PdfReader(in_path)
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
                with open(out_path, "w", encoding="utf-8") as f:
                    f.write(text)
            await asyncio.to_thread(extract_pdf_txt)
            return out_path
        elif target_ext == "images":
            zip_out = CONVERTED_DIR / f"pdf_imgs_{generate_temp_path('zip').name}"
            success = await asyncio.to_thread(convert_pdf_to_images_sync, in_path, zip_out)
            if success and zip_out.exists():
                return zip_out
            raise RuntimeError("PDF page extraction failed.")
            
    # 5. Documents / Office conversions via LibreOffice Headless
    if source_ext in ("docx", "odt", "rtf", "pdf", "html", "txt", "xlsx", "ods", "pptx", "odp"):
        # LibreOffice supports almost all these transitions (including docx to pdf, epub, etc.)
        try:
            temp_out = await asyncio.to_thread(convert_document_libreoffice_sync, in_path, target_ext, CONVERTED_DIR)
            shutil.move(temp_out, out_path)
            return out_path
        except Exception as e:
            logger.error(f"LibreOffice failed: {e}. Trying fallback...")
            if source_ext == "docx":
                if target_ext == "txt":
                    def docx_to_txt():
                        text = docx2txt.process(str(in_path))
                        out_path.write_text(text, encoding="utf-8")
                    await asyncio.to_thread(docx_to_txt)
                    return out_path
                elif target_ext == "html":
                    def docx_to_html():
                        with open(in_path, "rb") as f:
                            result = mammoth.convert_to_html(f)
                            out_path.write_text(result.value, encoding="utf-8")
                    await asyncio.to_thread(docx_to_html)
                    return out_path
            raise e
            
    # 6. Audio / Video conversions via FFmpeg
    if source_ext in ("mp4", "avi", "mkv", "mov", "webm", "mp3", "wav", "ogg", "aac"):
        await asyncio.to_thread(convert_media_ffmpeg_sync, in_path, out_path, target_ext)
        return out_path
        
    raise NotImplementedError(f"Unsupported conversion path from {source_ext} to {target_ext}.")
