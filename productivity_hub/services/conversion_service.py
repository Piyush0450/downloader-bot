"""
Smart Document, Media & File Conversion Service.

Executes file format conversions across Images, Audio, Video, Documents, Spreadsheets, and Archives.
Integrates magic-byte signature file detection, thread-pool execution for CPU tasks, subprocess
execution for FFmpeg and LibreOffice with strict timeouts, safe filename generation, and an
automated circuit breaker to temporarily disable repeatedly failing conversions.
"""

import asyncio
import csv
import io
import json
import logging
import shutil
import subprocess
import tarfile
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from PIL import Image
import pypdf

from productivity_hub.config.settings import Settings, get_settings
from productivity_hub.utils.file_detector import detect_file_type
from productivity_hub.utils.conversion_registry import get_conversion_registry, ConversionRule
from productivity_hub.utils.safe_filename import generate_safe_filepath

logger = logging.getLogger("productivity_hub.services.conversion")

CONVERTED_DIR = Path("temp/converted")


class ConversionError(Exception):
    """Raised when file conversion fails or backend is temporarily unavailable."""
    pass


@dataclass
class CircuitBreakerStatus:
    """Tracks consecutive failure count and cooldown timestamp for a conversion pair."""
    failure_count: int = 0
    disabled_until: float = 0.0


class ConversionService:
    """
    Async conversion engine executing document, image, audio, video, and spreadsheet conversions.
    """

    def __init__(self, settings: Optional[Settings] = None):
        self.settings = settings or get_settings()
        CONVERTED_DIR.mkdir(parents=True, exist_ok=True)
        # Map: (source_ext, target_ext) -> CircuitBreakerStatus
        self._circuit_breakers: Dict[Tuple[str, str], CircuitBreakerStatus] = {}

    def is_conversion_available(self, source_format: str, target_format: str) -> bool:
        """
        Check if conversion is supported by registry AND not disabled by circuit breaker.
        """
        registry = get_conversion_registry()
        if not registry.is_available(source_format, target_format):
            return False

        key = (source_format.lower(), target_format.lower())
        status = self._circuit_breakers.get(key)
        if status and time.time() < status.disabled_until:
            logger.warning(f"Conversion {source_format}->{target_format} is temporarily disabled by Circuit Breaker.")
            return False

        return True

    async def convert_file(
        self,
        source_path: Path,
        target_format: str,
        user_id: int
    ) -> Dict[str, Any]:
        """
        Detect file type, validate conversion rule, and execute conversion asynchronously.

        Args:
            source_path: Local Path of input file.
            target_format: Extension string of requested output format ('pdf', 'jpg', 'mp3', etc.).
            user_id: Telegram numeric user ID.

        Returns:
            Dict[str, Any]: Conversion metadata (result_file, file_size, source_format, target_format).

        Raises:
            ConversionError: If file is unsupported, library missing, or conversion fails.
        """
        if not source_path.exists() or not source_path.is_file():
            raise ConversionError("Input file does not exist or is unreadable.")

        # 1. Smart File Type Detection
        file_info = detect_file_type(source_path)
        source_format = file_info.extension.lower()
        target_format = target_format.strip().lower().lstrip(".")

        logger.info(f"Initiating conversion for user {user_id}: {source_path.name} ({source_format} -> {target_format})")

        # 2. Check Circuit Breaker & Registry
        if not self.is_conversion_available(source_format, target_format):
            raise ConversionError(f"Conversion from .{source_format} to .{target_format} is currently unavailable.")

        registry = get_conversion_registry()
        rule = registry.get_rule(source_format, target_format)
        timeout = rule.timeout_seconds if rule else 60

        key = (source_format, target_format)
        target_path, _ = generate_safe_filepath(CONVERTED_DIR, f"conv_{source_path.stem}.{target_format}", prefix="conv")

        # 3. Execute Conversion in Thread Pool
        loop = asyncio.get_running_loop()
        try:
            result_path = await asyncio.wait_for(
                loop.run_in_executor(
                    None, self._sync_execute_conversion, source_path, target_path, source_format, target_format
                ),
                timeout=timeout
            )

            # Validate that generated output exists and is non-empty
            if not result_path.exists() or result_path.stat().st_size == 0:
                raise ConversionError(f"Generated conversion file is empty or missing.")

            # Reset Circuit Breaker on success
            if key in self._circuit_breakers:
                self._circuit_breakers[key].failure_count = 0

        except asyncio.TimeoutError:
            status = self._circuit_breakers.setdefault(key, CircuitBreakerStatus())
            status.failure_count += 1
            logger.error(f"Conversion timed out after {timeout}s for {source_path.name} ({source_format}->{target_format})")
            raise ConversionError(f"Conversion timed out after {timeout} seconds.")

        except Exception as e:
            status = self._circuit_breakers.setdefault(key, CircuitBreakerStatus())
            status.failure_count += 1
            if status.failure_count >= 3:
                status.disabled_until = time.time() + 900  # Disable for 15 minutes
                logger.error(f"CIRCUIT BREAKER TRIGGERED: Disabling {source_format}->{target_format} for 15m after 3 consecutive failures.")

            logger.error(f"Conversion failed for {source_path.name} ({source_format}->{target_format}): {e}", exc_info=True)
            raise ConversionError(f"File conversion failed: {str(e)}")

        file_size_bytes = result_path.stat().st_size
        return {
            "result_file": result_path,
            "file_name": result_path.name,
            "file_size": file_size_bytes,
            "file_size_mb": file_size_bytes / (1024 * 1024),
            "source_format": source_format,
            "target_format": target_format,
        }

    def _sync_execute_conversion(
        self,
        source_path: Path,
        target_path: Path,
        source_format: str,
        target_format: str
    ) -> Path:
        """Synchronous dispatcher executing format conversion routines."""

        # ─── 1. Images (Pillow) ──────────────────────────────────────────
        image_exts = ("png", "jpg", "jpeg", "webp", "bmp", "tiff", "ico")
        if source_format in image_exts and target_format in image_exts:
            with Image.open(source_path) as img:
                fmt_str = "JPEG" if target_format in ("jpg", "jpeg") else target_format.upper()
                if target_format in ("jpg", "jpeg") and img.mode in ("RGBA", "LA", "P"):
                    rgb_img = Image.new("RGB", img.size, (255, 255, 255))
                    rgb_img.paste(img, mask=img.split()[-1] if "A" in img.mode else None)
                    rgb_img.save(target_path, format=fmt_str, quality=90)
                elif target_format == "ico":
                    img.save(target_path, format="ICO", sizes=[(32, 32), (64, 64)])
                else:
                    img.save(target_path, format=fmt_str)
            return target_path

        # Image to PDF
        if source_format in image_exts and target_format == "pdf":
            with Image.open(source_path) as img:
                rgb_img = img.convert("RGB")
                rgb_img.save(target_path, "PDF", resolution=100.0)
            return target_path

        # ─── 2. Audio & Video via FFmpeg ──────────────────────────────────
        audio_exts = ("mp3", "wav", "aac", "ogg", "flac", "m4a", "opus")
        video_exts = ("mp4", "mkv", "webm", "avi", "mov")

        if (source_format in audio_exts and target_format in audio_exts) or \
           (source_format in video_exts and target_format in video_exts) or \
           (source_format in video_exts and target_format in audio_exts):

            cmd = ["ffmpeg", "-y", "-i", str(source_path)]
            if target_format == "mp3":
                cmd.extend(["-vn", "-acodec", "libmp3lame", "-b:a", "320k"])
            elif target_format == "wav":
                cmd.extend(["-vn", "-acodec", "pcm_s16le"])
            elif target_format == "aac":
                cmd.extend(["-vn", "-acodec", "aac", "-b:a", "256k"])
            elif target_format == "flac":
                cmd.extend(["-vn", "-acodec", "flac"])
            elif target_format == "ogg":
                cmd.extend(["-vn", "-acodec", "libvorbis"])
            elif target_format == "opus":
                cmd.extend(["-vn", "-acodec", "libopus", "-b:a", "128k"])
            elif target_format == "m4a":
                cmd.extend(["-vn", "-acodec", "aac", "-b:a", "256k"])
            elif target_format == "mp4":
                cmd.extend(["-c:v", "libx264", "-c:a", "aac"])
            elif target_format == "webm":
                cmd.extend(["-c:v", "libvpx-vp9", "-c:a", "libopus"])
            elif target_format == "mkv":
                cmd.extend(["-c:v", "copy", "-c:a", "copy"])

            cmd.append(str(target_path))
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
            if res.returncode != 0:
                err = res.stderr.decode("utf-8", errors="replace")
                raise ConversionError(f"FFmpeg error ({res.returncode}): {err[-300:]}")
            return target_path

        # ─── 3. Documents ────────────────────────────────────────────────
        # PDF -> TXT
        if source_format == "pdf" and target_format == "txt":
            reader = pypdf.PdfReader(str(source_path))
            extracted_text = []
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                extracted_text.append(f"--- Page {idx + 1} ---\n{text}")
            target_path.write_text("\n\n".join(extracted_text), encoding="utf-8")
            return target_path

        # TXT -> PDF
        if source_format == "txt" and target_format == "pdf":
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas

            raw_text = source_path.read_text(encoding="utf-8", errors="replace")
            c = canvas.Canvas(str(target_path), pagesize=letter)
            width, height = letter
            y = height - 40

            for line in raw_text.splitlines():
                if y < 40:
                    c.showPage()
                    y = height - 40
                c.drawString(40, y, line[:100])
                y -= 15

            c.save()
            return target_path

        # DOCX -> TXT
        if source_format == "docx" and target_format == "txt":
            import docx
            doc = docx.Document(str(source_path))
            paragraphs = [p.text for p in doc.paragraphs]
            target_path.write_text("\n\n".join(paragraphs), encoding="utf-8")
            return target_path

        # LibreOffice conversions (DOCX/DOC/ODT/RTF/PPTX/PPT -> PDF)
        office_formats = ("docx", "doc", "odt", "rtf", "pptx", "ppt")
        if source_format in office_formats and target_format == "pdf":
            soffice_cmd = shutil.which("soffice") or shutil.which("libreoffice")
            if not soffice_cmd:
                raise ConversionError("LibreOffice executable not found on system PATH.")

            out_dir = target_path.parent
            res = subprocess.run(
                [soffice_cmd, "--headless", "--convert-to", "pdf", str(source_path), "--outdir", str(out_dir)],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=120
            )
            expected_output = out_dir / f"{source_path.stem}.pdf"
            if expected_output.exists():
                if expected_output != target_path:
                    shutil.move(str(expected_output), str(target_path))
                return target_path
            raise ConversionError(f"LibreOffice conversion failed: {res.stderr.decode('utf-8', errors='replace')}")

        # MD / HTML -> TXT
        if source_format in ("md", "html") and target_format == "txt":
            import re
            content = source_path.read_text(encoding="utf-8", errors="replace")
            clean_text = re.sub(r"<[^>]+>", "", content)
            target_path.write_text(clean_text, encoding="utf-8")
            return target_path

        # ─── 4. Spreadsheets ─────────────────────────────────────────────
        # XLSX -> CSV / TSV / JSON
        if source_format == "xlsx" and target_format in ("csv", "tsv", "json"):
            import openpyxl
            wb = openpyxl.load_workbook(source_path, data_only=True)
            sheet = wb.active
            delimiter = "\t" if target_format == "tsv" else ","

            if target_format in ("csv", "tsv"):
                with open(target_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f, delimiter=delimiter)
                    for row in sheet.iter_rows(values_only=True):
                        writer.writerow(row)
                wb.close()
                return target_path
            elif target_format == "json":
                rows = list(sheet.iter_rows(values_only=True))
                data = []
                if rows:
                    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
                    for r in rows[1:]:
                        data.append(dict(zip(headers, r)))
                wb.close()
                target_path.write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")
                return target_path

        # CSV -> XLSX / TSV / JSON
        if source_format == "csv" and target_format in ("xlsx", "tsv", "json"):
            with open(source_path, "r", encoding="utf-8", errors="replace") as f:
                reader = list(csv.reader(f))

            if target_format == "xlsx":
                import openpyxl
                wb = openpyxl.Workbook()
                ws = wb.active
                for r in reader:
                    ws.append(r)
                wb.save(target_path)
                wb.close()
                return target_path
            elif target_format == "tsv":
                with open(target_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f, delimiter="\t")
                    for r in reader:
                        writer.writerow(r)
                return target_path
            elif target_format == "json":
                data = []
                if reader:
                    headers = reader[0]
                    for r in reader[1:]:
                        data.append(dict(zip(headers, r)))
                target_path.write_text(json.dumps(data, indent=2), encoding="utf-8")
                return target_path

        # ─── 5. Archives ─────────────────────────────────────────────────
        # TAR -> ZIP
        if source_format == "tar" and target_format == "zip":
            with tarfile.open(source_path, "r") as tf, zipfile.ZipFile(target_path, "w", zipfile.ZIP_DEFLATED) as zf:
                for member in tf.getmembers():
                    if member.isfile():
                        f = tf.extractfile(member)
                        if f:
                            zf.writestr(member.name, f.read())
            return target_path

        # ZIP -> TAR
        if source_format == "zip" and target_format == "tar":
            with zipfile.ZipFile(source_path, "r") as zf, tarfile.open(target_path, "w") as tf:
                for name in zf.namelist():
                    info = zf.getinfo(name)
                    if not info.is_dir():
                        data = zf.read(name)
                        ti = tarfile.TarInfo(name=name)
                        ti.size = len(data)
                        tf.addfile(ti, io.BytesIO(data))
            return target_path

        raise ConversionError(f"No active conversion handler implemented for {source_format} -> {target_format}")


# Global ConversionService singleton
_conversion_service_instance: Optional[ConversionService] = None


def get_conversion_service() -> ConversionService:
    """Retrieve global ConversionService instance."""
    global _conversion_service_instance
    if _conversion_service_instance is None:
        _conversion_service_instance = ConversionService()
    return _conversion_service_instance
